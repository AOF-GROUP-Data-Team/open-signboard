# ============================================================
# Signboard Report — Zenput 600319 (Classic + Lubda) + 671663 (Garatis)
# GitHub Actions. Today's submissions (KSA): signboard photo +
# "Is the signboard working properly?" per branch, plus branches that
# did not submit -> xlsx, emailed as attachment with an Arabic summary.
# Secrets: ZENPUT_API_TOKEN, GMAIL_APP_PASSWORD
# Env: MODE = auto (send) | test (no email), REPORT_DATE = "" or YYYY-MM-DD
# ============================================================
import os, sys, smtplib, requests, re, io
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email.utils import formataddr
from email import encoders
from datetime import datetime
import pytz
from PIL import Image as PILImage, ImageOps
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

API_TOKEN   = os.environ.get("ZENPUT_API_TOKEN", "")
GMAIL_USER  = "aof.group.auto@gmail.com"
GMAIL_PASS  = os.environ.get("GMAIL_APP_PASSWORD", "")
SENDER_NAME = "Business Intelligence"
TO_EMAILS   = ["o.salahaddin@aofgroup.com"]
CC_EMAILS   = ["a.alsalem@aofgroup.com"]
MODE        = os.environ.get("MODE", "auto").strip().lower()
REPORT_DATE = os.environ.get("REPORT_DATE", "").strip()   # "" = today (KSA)
BASE        = "https://www.zenput.com"
HEADERS     = {"X-API-TOKEN": API_TOKEN}
TZ          = pytz.timezone("Asia/Riyadh")
PAGE        = 100
LINK_BASE   = "https://ops-execution.crunchtime.com/reports/#form_id/"
PHOTO_W     = 260         # px width of embedded photo
ROW_H       = 150         # minimum row height (points) for photo rows

# template -> brand
TEMPLATES = {600319: "Shawarma Classic", 671663: "Garatis"}

# Expected branches per brand: code -> (Zenput location ID or None, name). B03 deprecated.
BRANCHES = {
    "Shawarma Classic": {
        "B01": (2155652, "NURUH"),  "B02": (2164013, "KHRUH"),  "B04": (2197297, "NSRUH"),
        "B05": (2164016, "RWRUH"),  "B06": (2164017, "DARUH"),  "B07": (2197299, "LBRUH"),
        "B08": (2164019, "SWRUH"),  "B09": (2164020, "AZRUH"),  "B10": (2164021, "SHRUH"),
        "B11": (2164022, "NRRUH"),  "B12": (2164023, "TWRUH"),  "B13": (2155654, "AQRUH"),
        "B14": (2164025, "RBRUH"),  "B15": (2164026, "NDRUH"),  "B16": (2164027, "BDRUH"),
        "B17": (2164028, "QRRUH"),  "B18": (2197298, "TKRUH"),  "B19": (2164030, "MURUH"),
        "B21": (2164032, "KRRUH"),  "B22": (2185452, "OBJED"),  "B23": (2190657, "SLAHS"),
        "B24": (2164031, "SFJED"),  "B25": (2199002, "RWAHS"),  "B26": (2199835, "HAJED"),
        "B27": (2203271, "SARUH"),  "B28": (2210205, "MAJED"),  "B30": (2211854, "QARUH"),
        "B31": (2235670, "ANRUH"),  "B32": (2239240, "FYJED"),  "B33": (2242934, "HIRJED"),
        "B34": (2243963, "URRUH"),  "B35": (2250799, "IRRUH"),  "B36": (2258220, "PSJED"),
        "B37": (2257790, "SHWMAK"), "B38": (2260889, "UHDMM"),  "B39": (2263062, "HSRUH"),
        "B40": (2281339, "MZDMM"),
        "LB01": (2169459, "Lubda Alaqeq"), "LB02": (2222802, "Lubda Alkhaleej"),
    },
    "Garatis": {
        "QB01": (2232755, "السويدي"), "QB02": (2235805, "النرجس"), "QB03": (2254072, "اليرموك"),
        "QB04": (2256386, "العقيق"),  "QB05": (2268360, "الخافجي"), "QB06": (2270650, "الخليج"),
        "QB07": (2274188, "الندوة"),  "QB08": (2276794, "المنصورة"), "QB09": (None, "النزهة"),
    },
}

# ---------------- helpers ----------------
def parse_dt(s):
    if not s: return None
    try: return datetime.fromisoformat(s.replace("Z", "+00:00")).astimezone(TZ)
    except Exception: return None

def resolve_branch(brand, loc):
    exp = BRANCHES[brand]
    try: lid = int(loc.get("id"))
    except Exception: lid = None
    for code, (i, _) in exp.items():
        if i is not None and i == lid: return code
    name = re.sub(r"\s+", " ", (loc.get("name") or loc.get("display_name") or "")).strip()
    if brand == "Garatis":
        m = re.search(r"\bQB\s?(\d{2})\b", name, re.I) or re.search(r"Garatis+\s*B(\d{2})", name, re.I)
        code = f"QB{m.group(1)}" if m else None
        return code if code in exp else None
    m = re.match(r"([A-Z]+)\s+(B\d{2})\b", name.upper())
    if m and m.group(2) in exp and exp[m.group(2)][1] == m.group(1):
        return m.group(2)
    for code, (_, n) in exp.items():
        if code.startswith("LB") and n.lower() in name.lower(): return code
    return None

def is_photo_q(t):   return "signboard" in t and "open" in t and "drive" not in t
def is_working_q(t): return "signboard" in t and "working" in t

def photo_keys(v):
    items = v if isinstance(v, list) else ([v] if isinstance(v, dict) else [])
    return [i["s3_key"] for i in items if isinstance(i, dict) and i.get("s3_key")]

def signed_url(s3_key):
    r = requests.get(f"{BASE}/api/v2/users/current/storage/", headers=HEADERS,
                     params={"path": s3_key}, timeout=60)
    if r.status_code != 200: return None
    return ((r.json() or {}).get("data") or {}).get("location")

def fetch_photo(s3_key):
    url = signed_url(s3_key)
    if not url: return None, "no signed url"
    r = requests.get(url, timeout=60)
    if r.status_code != 200: return None, f"HTTP {r.status_code}"
    try:
        img = PILImage.open(io.BytesIO(r.content))
        try: img = ImageOps.exif_transpose(img)
        except Exception: pass
        img = img.convert("RGB")
        img = img.resize((PHOTO_W, int(img.height * PHOTO_W / img.width)))
        buf = io.BytesIO(); img.save(buf, "JPEG", quality=80); buf.seek(0)
        return buf, img.height
    except Exception as e:
        return None, str(e)

def fetch_day(template_id, day):
    out, start = [], 0
    while True:
        r = requests.get(f"{BASE}/api/v3/submissions/", headers=HEADERS, params={
            "form_template_id": template_id, "limit": PAGE, "start": start}, timeout=60)
        print(f"[{template_id}] GET start={start} -> HTTP {r.status_code}")
        if r.status_code != 200:
            print(r.text[:500]); break
        data = r.json().get("data", [])
        if not data: break
        out.extend(data)
        dts = [d for d in (parse_dt((s.get("smetadata") or {}).get("date_submitted")) for s in data) if d]
        if len(data) < PAGE or (dts and min(dts).date() < day): break
        start += PAGE
    return out

if not API_TOKEN:
    sys.exit("ZENPUT_API_TOKEN is not set")

# ---------------- fetch + latest per branch ----------------
now = datetime.now(TZ)
day = datetime.strptime(REPORT_DATE, "%Y-%m-%d").date() if REPORT_DATE else now.date()
print(f"Report date: {day} (KSA)")

rows = []
for tid, brand in TEMPLATES.items():
    subs = fetch_day(tid, day)
    latest, unknown = {}, set()
    for s in subs:
        sm = s.get("smetadata") or {}
        dt = parse_dt(sm.get("date_submitted"))
        if not dt or dt.date() != day: continue
        loc = sm.get("location") or {}
        code = resolve_branch(brand, loc)
        if code is None:
            unknown.add(f"{loc.get('id')} {loc.get('name')}"); continue
        if code not in latest or dt > latest[code][0]:
            latest[code] = (dt, s)
    print(f"[{tid}] {brand}: fetched {len(subs)} | submitted on {day}: {len(latest)} branches")
    if unknown: print(f"!! [{tid}] unmapped locations (ignored): {unknown}")

    warned = False
    for code, (_, name) in sorted(BRANCHES[brand].items(), key=lambda x: (x[0].startswith("LB"), int(re.sub(r"\D", "", x[0])))):
        if code not in latest:
            rows.append({"brand": brand, "code": code, "name": name, "submitted": False}); continue
        dt, s = latest[code]
        working, keys, prepared, found = None, [], "", set()
        for a in s.get("answers", []):
            t = (a.get("title") or "").strip().lower()
            if a.get("field_type") == "image" and is_photo_q(t):
                keys = photo_keys(a.get("value")); found.add("photo")
            elif is_working_q(t):
                working = str(a.get("value")).lower() in ("true", "yes", "1"); found.add("working")
            elif t == "prepared by":
                prepared = str(a.get("value") or "").strip()
        if len(found) < 2 and not warned:
            warned = True
            print(f"\n!! [{tid}] signboard questions not found in {code} — answer titles:")
            for a in s.get("answers", []):
                print(f"   {a.get('title')} / {a.get('field_type')}")
        rows.append({"brand": brand, "code": code, "name": name, "submitted": True,
                     "time": dt.strftime("%H:%M"), "prepared": prepared, "working": working,
                     "keys": keys, "link": LINK_BASE + str(s.get("legacy_submission_id") or "")})

# ---------------- xlsx ----------------
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
RED = PatternFill("solid", fgColor="F8CBAD")
GRN = PatternFill("solid", fgColor="C6EFCE")
GRY = PatternFill("solid", fgColor="D9D9D9")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook(); ws = wb.active; ws.title = "Signboard"
ws.append(["Branch", "Status", "Time Submitted", "Prepared By", "Signboard Working?", "Signboard Photo", "Link"])
for c in ws[1]: c.fill, c.font, c.alignment = HDR, Font(bold=True, color="FFFFFF"), CENTER
for col, w in zip("ABCDEFG", [22, 15, 14, 16, 16, 38, 9]): ws.column_dimensions[col].width = w

photo_fail = 0
for brand in TEMPLATES.values():
    ws.append([brand]); i = ws.max_row
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
    ws[f"A{i}"].fill, ws[f"A{i}"].font, ws[f"A{i}"].alignment = SEC, Font(bold=True, size=12), CENTER
    for r in [x for x in rows if x["brand"] == brand]:
        label = f"{r['code']} {r['name']}"
        if not r["submitted"]:
            ws.append([label, "Not submitted", "", "", "", "", ""]); i = ws.max_row
            for c in ws[i]: c.fill = GRY
            ws[f"B{i}"].fill, ws[f"B{i}"].font = RED, Font(bold=True, color="C00000")
        else:
            ws.append([label, "Submitted", r["time"], r["prepared"],
                       "Yes" if r["working"] else ("No" if r["working"] is False else "-"), "", "Open"])
            i = ws.max_row
            ws[f"E{i}"].fill = GRN if r["working"] else RED
            ws[f"G{i}"].hyperlink, ws[f"G{i}"].font = r["link"], Font(color="0563C1", underline="single")
            if r["keys"]:
                buf, info = fetch_photo(r["keys"][0])
                if buf:
                    ws.add_image(XLImage(buf), f"F{i}")
                    ws.row_dimensions[i].height = max(ROW_H, info * 0.75 + 6)
                else:
                    photo_fail += 1; ws[f"F{i}"] = f"photo failed: {info}"
                    print(f"  photo failed {label}: {info}")
            else:
                ws[f"F{i}"] = "No photo"; ws[f"F{i}"].fill = RED
        for c in ws[ws.max_row]: c.border, c.alignment = THIN, CENTER

ws.append([])
for brand in TEMPLATES.values():
    br = [r for r in rows if r["brand"] == brand]
    ws.append([f"{brand}: submitted {sum(1 for r in br if r['submitted'])} / {len(br)} | "
               f"not working {sum(1 for r in br if r['submitted'] and r['working'] is False)}"])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
ws.freeze_panes = "A2"

fname = f"signboard_report_{day}.xlsx"
wb.save(fname)
print(f"\nSaved {fname} | photo failures: {photo_fail}")
for brand in TEMPLATES.values():
    br = [r for r in rows if r["brand"] == brand]
    print(f"\n{brand}: submitted {sum(1 for r in br if r['submitted'])}/{len(br)}")
    print("  Not working:", [f"{r['code']} {r['name']}" for r in br if r["submitted"] and r["working"] is False] or "-")
    print("  Not submitted:", [f"{r['code']} {r['name']}" for r in br if not r["submitted"]] or "-")


# ---------------- email ----------------
BRAND_AR = {"Shawarma Classic": "شاورما كلاسيك", "Garatis": "قراطيس"}
TD = "padding:6px 12px;border:1px solid #cccccc;text-align:center"
TH = "padding:8px 12px;background:#1F3864;color:#ffffff;text-align:center"

def list_table(title, items):
    if not items:
        return f"<p style='margin:6px 0'><b>{title}:</b> لا يوجد</p>"
    body = "".join(f"<tr><td style='{TD}'>{i}</td><td style='{TD};color:#C00000' dir='auto'>{b}</td></tr>"
                   for i, b in enumerate(items, 1))
    return (f"<p style='margin:10px 0 4px'><b>{title} ({len(items)})</b></p>"
            f"<table style='border-collapse:collapse'><tr><th style='{TH}'>#</th><th style='{TH}'>الفرع</th></tr>{body}</table>")

def build_html():
    summary, details = "", ""
    for brand in TEMPLATES.values():
        br = [r for r in rows if r["brand"] == brand]
        sub = sum(1 for r in br if r["submitted"])
        ok = [r for r in br if r["submitted"] and r["working"]]
        bad = [f"{r['code']} {r['name']}" for r in br if r["submitted"] and r["working"] is False]
        miss = [f"{r['code']} {r['name']}" for r in br if not r["submitted"]]
        ar = BRAND_AR.get(brand, brand)
        summary += (f"<tr><td style='{TD}'><b>{ar}</b></td><td style='{TD}'>{sub} / {len(br)}</td>"
                    f"<td style='{TD};color:#006100'><b>{len(ok)}</b></td>"
                    f"<td style='{TD};color:#C00000'><b>{len(bad)}</b></td>"
                    f"<td style='{TD};color:#C00000'><b>{len(miss)}</b></td></tr>")
        details += (f"<h3 style='color:#1F3864;margin:20px 0 6px'>{ar}</h3>"
                    + list_table("فروع اللوحة لا تعمل", bad)
                    + list_table("فروع لم تسلّم النموذج", miss))
    return f"""<html><body dir="rtl" style="font-family:Tahoma,Arial,sans-serif;font-size:14px;color:#222;text-align:right">
<p>السلام عليكم ورحمة الله وبركاته،</p>
<p>مرفق لكم تقرير لوحات الفروع ليوم <span dir="ltr">{day}</span>، ويشمل صورة اللوحة وحالة تشغيلها لكل فرع.</p>
<table style="border-collapse:collapse">
<tr><th style="{TH}">العلامة التجارية</th><th style="{TH}">الفروع المسلّمة</th><th style="{TH}">اللوحات تعمل</th>
<th style="{TH}">اللوحات لا تعمل</th><th style="{TH}">فروع لم تسلّم</th></tr>
{summary}
</table>
{details}
<p style="margin-top:18px">صور اللوحات وروابط التسليم في الملف المرفق.</p>
<p>مع خالص التحية،<br>إدارة ذكاء الأعمال</p>
</body></html>"""

def send_email(path):
    if not GMAIL_PASS:
        sys.exit("GMAIL_APP_PASSWORD is not set")
    msg = MIMEMultipart()
    msg["From"] = formataddr((SENDER_NAME, GMAIL_USER))
    msg["To"] = ", ".join(TO_EMAILS)
    msg["Cc"] = ", ".join(CC_EMAILS)
    msg["Subject"] = f"تقرير لوحات الفروع - {day}"
    msg.attach(MIMEText(build_html(), "html", "utf-8"))
    with open(path, "rb") as f:
        part = MIMEBase("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{os.path.basename(path)}"')
    msg.attach(part)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as s:
        s.login(GMAIL_USER, GMAIL_PASS)
        s.sendmail(GMAIL_USER, TO_EMAILS + CC_EMAILS, msg.as_string())
    print(f"Email sent to {TO_EMAILS} cc {CC_EMAILS}")

if MODE == "test":
    print("MODE=test -> email skipped")
else:
    send_email(fname)
